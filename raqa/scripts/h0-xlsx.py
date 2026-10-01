#!/usr/bin/env python3
"""Studi H0 dalam format Excel (rater kedua untuk klasifikasi 19 kegagalan baseline, Tabel 4 naskah).

  python3 scripts/h0-xlsx.py build                          -> paket-reviewer/H0-Klasifikasi-Kegagalan.xlsx
  python3 scripts/h0-xlsx.py import <file.xlsx> [--inisial X] -> review/h0-sheet-<inisial>.csv
  node scripts/h0-failure-classification.js analyze review/h0-sheet-<inisial>.csv

Sumber item: review/h0-sheet.csv (dibuat oleh `h0-failure-classification.js template`). Label penulis ada di
review/h0-KEY-jangan-dibagikan.json dan TIDAK dimasukkan ke Excel.
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAQA = Path(__file__).resolve().parent.parent
E2E = RAQA.parent
SHEET = RAQA / 'review' / 'h0-sheet.csv'
OUT = RAQA / 'paket-reviewer' / 'H0-Klasifikasi-Kegagalan.xlsx'
CATS = ['product', 'test', 'timing', 'environment']
CLASS_COL = 'class_product_test_timing_environment'

FONT = 'Arial'
F = Font(name=FONT, size=10)
FB = Font(name=FONT, size=10, bold=True)
FH = Font(name=FONT, size=10, bold=True, color='FFFFFF')
FT = Font(name=FONT, size=14, bold=True)
HEAD = PatternFill('solid', fgColor='305496')
INPUT = PatternFill('solid', fgColor='FFF2CC')
GREY = PatternFill('solid', fgColor='F2F2F2')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def cell(ws, r, c, v, font=F, fill=None):
    x = ws.cell(row=r, column=c, value=v)
    x.font, x.alignment, x.border = font, WRAP, BOX
    if fill:
        x.fill = fill
    return x


def para(ws, r, text, font=F, cols=8):
    x = ws.cell(row=r, column=1, value=text)
    x.font, x.alignment = font, Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
    if font is F:
        ws.row_dimensions[r].height = max(15, 15 * (len(text) // 110 + 1))
    return r + 1


def dropdown(ws, formula, rng, msg):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorTitle='Nilai tidak valid', error=msg, promptTitle='Pilih', prompt=msg, showInputMessage=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def build():
    with open(SHEET, encoding='utf-8-sig', newline='') as f:
        items = list(csv.DictReader(f))
    catalog = {s['id']: s for s in json.loads((E2E / 'scenarios' / 'catalog.json').read_text())['scenarios']}

    wb = Workbook()
    ws = wb.active
    ws.title = 'Petunjuk'
    ws.sheet_view.showGridLines = False
    for col in 'ABCDEFGH':
        ws.column_dimensions[col].width = 16
    ws.cell(row=1, column=1, value='Klasifikasi Kegagalan Test Otomatis — SIMASIS (H0)').font = FT
    for r, (label, hint) in enumerate([('Inisial rater', 'mis. AB'), ('Tanggal', 'dd-mm-yyyy')], start=3):
        ws.cell(row=r, column=1, value=label).font = FB
        c = ws.cell(row=r, column=2)
        c.fill, c.border = INPUT, BOX
        ws.cell(row=r, column=3, value=hint).font = Font(name=FONT, size=9, italic=True, color='808080')
    r = para(ws, 6, 'Persetujuan (informed consent)', FB)
    r = para(ws, r, '"Saya bersedia mengklasifikasikan kegagalan test otomatis untuk penelitian tesis Hilman Firdaus. Saya memahami '
                    'bahwa keikutsertaan saya sukarela, dapat saya hentikan kapan saja, nama saya tidak dicantumkan tanpa izin, dan '
                    'bahan ini tidak berisi data pribadi santri, asatidz, atau orang tua."')
    ws.cell(row=r, column=1, value='Saya setuju').font = FB  # baris 8
    agree = ws.cell(row=r, column=2)
    agree.fill, agree.border = INPUT, BOX
    dropdown(ws, '"ya,tidak"', f'B{r}', 'ya atau tidak')
    r += 2
    for text, font in [
        ('Tugas Anda', FB),
        ('Suite test otomatis (Playwright) dijalankan 10 kali terhadap portal SIMASIS. Sheet "Kegagalan" berisi 19 test yang '
         'gagal. Untuk setiap baris, tentukan PENYEBAB kegagalannya dengan memilih satu kategori di kolom kuning. Tidak ada '
         'jawaban yang "diharapkan"; nilai sesuai penilaian profesional Anda. Perkiraan waktu: ±30 menit.', F),
        ('', F),
        ('Kategori', FB),
        ('product — cacat pada APLIKASI: aplikasi berperilaku tidak sesuai acceptance criteria (test-nya benar, aplikasinya salah).', F),
        ('test — cacat pada KODE TEST: locator salah, ekspektasi keliru, atau test menguji hal yang tidak ada (aplikasinya benar).', F),
        ('timing — hasil BERUBAH-UBAH antar-run karena waktu (animasi, loading, race condition), bukan karena kode yang salah.', F),
        ('environment — masalah SERVER/INFRASTRUKTUR (jaringan, server lambat, akses ditolak), bukan kode aplikasi atau kode test.', F),
        ('', F),
        ('Petunjuk membaca baris', FB),
        ('• verdict: stable-fail = gagal di semua 10 run; flaky = kadang lulus kadang gagal. passes = jumlah run yang lulus.', F),
        ('• project: chromium-desktop = browser desktop; mobile-chrome = tampilan HP.', F),
        ('• goal dan acceptance_criteria = apa yang SEHARUSNYA dipenuhi aplikasi. first_error = pesan error pertama dari test.', F),
        ('• Bila perlu, buka portal di browser untuk mengecek sendiri. Bila ragu, pilih kategori yang paling mungkin dan jelaskan di kolom komentar.', F),
        ('', F),
        ('Aturan', FB),
        ('1. Kerjakan sendiri, jangan berdiskusi dengan penulis atau rater lain sebelum selesai.  2. Hanya isi kolom kuning.  '
         '3. Simpan file sebagai H0-<inisial>.xlsx dan kirim kembali ke penulis.', F),
    ]:
        r = para(ws, r, text, font)

    k = wb.create_sheet('Kegagalan')
    cols = ['item_id', 'scenario_id', 'project', 'test_title', 'goal', 'acceptance_criteria', 'verdict', 'passes',
            'first_error', 'kategori', 'komentar']
    widths = [8, 12, 15, 34, 30, 34, 11, 8, 40, 14, 30]
    for i, (h, w) in enumerate(zip(cols, widths), start=1):
        x = cell(k, 1, i, h, FH, HEAD)
        k.column_dimensions[x.column_letter].width = w
    k.row_dimensions[1].height = 30
    for row, it in enumerate(items, start=2):
        sc = catalog.get(it['scenario_id'], {})
        crit = '\n'.join(f'{n}. {c}' for n, c in enumerate(sc.get('acceptance_criteria', []), start=1))
        vals = [it['item_id'], it['scenario_id'], it['project'], it['test_title'], sc.get('goal', ''), crit,
                it['verdict'], it['passes'], it['first_error']]
        for c, v in enumerate(vals, start=1):
            cell(k, row, c, v, fill=GREY)
        cell(k, row, 10, None, fill=INPUT)
        cell(k, row, 11, None, fill=INPUT)
        k.row_dimensions[row].height = 60
    last = len(items) + 1
    dropdown(k, '"' + ','.join(CATS) + '"', f'J2:J{last}', ' / '.join(CATS))
    k.freeze_panes = 'C2'

    p = wb.create_sheet('Progres')
    p.column_dimensions['A'].width = 30
    cell(p, 1, 1, 'Baris terisi', FB)
    cell(p, 1, 2, f'=COUNTA(Kegagalan!J2:J{last})')
    cell(p, 2, 1, 'Total baris', FB)
    cell(p, 2, 2, len(items))
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f'{OUT.relative_to(RAQA)} ({len(items)} item). Label penulis TIDAK disertakan.')


def import_xlsx(path, ini=None):
    wb = load_workbook(path, data_only=True)
    ini = ini or str(wb['Petunjuk']['B3'].value or '').strip()
    if not re.fullmatch(r'[A-Za-z0-9]{1,8}', ini):
        sys.exit('Inisial kosong/tidak valid (Petunjuk!B3); berikan --inisial.')
    if str(wb['Petunjuk']['B8'].value or '').strip().lower() != 'ya':
        print('PERINGATAN: persetujuan (consent) belum "ya".')
    ws = wb['Kegagalan']
    head = [str(c.value or '').strip() for c in ws[1]]
    rows = [dict(zip(head, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    rows = [r for r in rows if r.get('item_id')]
    bad = [f"{r['item_id']}: \"{r.get('kategori')}\"" for r in rows
           if str(r.get('kategori') or '').strip().lower() not in CATS]
    if bad:
        sys.exit('Kategori kosong/tidak valid (perbaiki bersama rater, jangan ditebak):\n  ' + '\n  '.join(bad))
    out = RAQA / 'review' / f'h0-sheet-{ini}.csv'
    # CSV minimal satu-baris-per-item: parser di h0-failure-classification.js membaca per baris,
    # sehingga komentar multi-baris diratakan.
    with open(out, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f)
        w.writerow(['item_id', CLASS_COL, 'comment'])
        for r in rows:
            w.writerow([r['item_id'], str(r['kategori']).strip().lower(), ' '.join(str(r.get('komentar') or '').split())])
    print(f'{len(rows)} item -> {out}')


if __name__ == '__main__':
    if len(sys.argv) >= 2 and sys.argv[1] == 'build':
        build()
    elif len(sys.argv) >= 3 and sys.argv[1] == 'import':
        ini = sys.argv[sys.argv.index('--inisial') + 1] if '--inisial' in sys.argv else None
        import_xlsx(sys.argv[2], ini)
    else:
        print(__doc__)
