#!/usr/bin/env python3
"""Excel diskusi (bahasa sehari-hari, tanpa istilah teknis) untuk item H0 yang tidak sepakat antara penulis
dan rater kedua.

  python3 scripts/h0-diskusi-xlsx.py build [review/h0-sheet-R2.csv]   -> review/H0-Diskusi-Ketidaksepakatan.xlsx
  python3 scripts/h0-diskusi-xlsx.py import <file.xlsx> [review/h0-sheet-R2.csv]
      -> review/h0-final-overrides.csv, lalu:
  node scripts/h0-failure-classification.js final review/h0-sheet-R2.csv review/h0-final-overrides.csv
"""
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

RAQA = Path(__file__).resolve().parent.parent
REVIEW = RAQA / 'review'
OUT = REVIEW / 'H0-Diskusi-Ketidaksepakatan.xlsx'

# Pilihan kesimpulan dalam bahasa biasa -> kategori analisis. Tanpa koma (pemisah daftar dropdown Excel).
CHOICES = {
    'Aplikasinya yang salah': 'product',
    'Cara pengecekan otomatisnya yang salah': 'test',
    'Hasilnya berubah-ubah tiap dicoba': 'timing',
    'Masalah server atau jaringan': 'environment',
}
PLAIN = {v: k for k, v in CHOICES.items()}

# Penjelasan tanpa istilah teknis, per (skenario, tampilan). Isi faktual diambil dari data baseline:
# jumlah percobaan dari results/metrics/baseline-local/per-test.csv, alasan penulis dari ground-truth.json.
TEXT = {
    ('S-AUTH-03', 'chromium-desktop'): {
        'dicek': 'Saat seseorang login dengan email dan password yang SALAH: (1) tetap berada di halaman login, '
                 '(2) muncul pesan kesalahan yang jelas, (3) kolom password dikosongkan.',
        'terjadi': 'Diperiksa otomatis 10 kali di tampilan laptop/komputer. Hasilnya 10 kali gagal: setelah login '
                   'dengan data salah, pesan kesalahan tidak muncul di layar.',
        'penulis': 'Aplikasinya yang salah: aplikasi tidak menampilkan pesan kesalahan ketika login gagal.',
        'tanya': '1. Saat Anda mencoba, email seperti apa yang dipakai? Pemeriksaan otomatis memakai email yang '
                 'penulisannya benar tetapi tidak terdaftar (contoh: orang.tidak.terdaftar@example.com), bukan '
                 'tulisan asal seperti "abc".\n'
                 '2. Apakah Anda mencoba di laptop/komputer, di alamat http://127.0.0.1:8000/login?\n'
                 '3. Mari coba bersama sekarang dengan email seperti contoh di atas: apakah pesan kesalahan muncul?',
    },
    ('S-HOME-05', 'mobile-chrome'): {
        'dicek': 'Di beranda, berita terbaru yang pertama bisa diklik, lalu membuka artikel yang judulnya sama '
                 'dengan judul di kartu berita.',
        'terjadi': 'Diperiksa otomatis 10 kali dalam TAMPILAN HP (layar kecil). Hasilnya 10 kali gagal: kartu '
                   'berita tidak bisa diklik. Catatan: pemeriksaan yang sama di tampilan laptop berhasil 10 dari 10.',
        'penulis': 'Aplikasinya yang salah: di tampilan HP, kartu berita tidak bisa diklik.',
        'tanya': '1. Anda mencobanya di HP (atau mode tampilan HP di browser), atau di laptop?\n'
                 '2. Mari coba bersama di HP: buka beranda, klik berita terbaru yang pertama. Apakah artikelnya '
                 'terbuka dan judulnya sama?',
    },
    ('S-QUAL-01', 'chromium-desktop'): {
        'dicek': 'Halaman daftar artikel santri (alamat /portal/artikel) tidak boleh punya masalah aksesibilitas '
                 'tingkat berat. Contoh masalah berat: tombol yang tidak punya nama, sehingga pengguna tunanetra '
                 '(yang memakai pembaca layar) tidak tahu fungsi tombol itu.',
        'terjadi': 'Diperiksa otomatis 10 kali. Hasilnya 10 kali gagal: pemeriksa aksesibilitas menemukan masalah '
                   'tingkat berat di halaman daftar artikel santri.',
        'penulis': 'Aplikasinya yang salah: tombol pencarian (ikon kaca pembesar) di halaman itu tidak diberi nama.',
        'tanya': '1. Halaman mana saja yang Anda periksa? Pemeriksaan otomatis memeriksa halaman DAFTAR artikel '
                 'santri (/portal/artikel), bukan halaman isi artikelnya.\n'
                 '2. Di halaman daftar itu, apakah tombol pencarian (ikon kaca pembesar) tercatat bermasalah?\n'
                 '3. Mari buka halaman daftar itu bersama dan periksa tombol pencariannya.',
    },
}

F = Font(name='Arial', size=11)
FB = Font(name='Arial', size=11, bold=True)
FT = Font(name='Arial', size=15, bold=True)
FI = Font(name='Arial', size=12, bold=True, color='FFFFFF')
HEAD = PatternFill('solid', fgColor='305496')
LABEL = PatternFill('solid', fgColor='F2F2F2')
INPUT = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')


def disagreements(rater_file):
    key = json.load(open(REVIEW / 'h0-KEY-jangan-dibagikan.json'))
    author = {k['itemId']: k['author'] for k in key}
    rater = {r['item_id']: r for r in csv.DictReader(open(rater_file, encoding='utf-8'))}
    sheet = {r['item_id']: r for r in csv.DictReader(open(REVIEW / 'h0-sheet.csv', encoding='utf-8-sig'))}
    col = 'class_product_test_timing_environment'
    return [(i, sheet[i], author[i], rater[i][col].strip().lower(), rater[i].get('comment', ''))
            for i in author if author[i] != rater[i][col].strip().lower()]


def row(ws, r, label, value, height=None, fill=None):
    a = ws.cell(row=r, column=1, value=label)
    a.font, a.fill, a.alignment, a.border = FB, LABEL, WRAP, BOX
    b = ws.cell(row=r, column=2, value=value)
    b.font, b.alignment, b.border = F, WRAP, BOX
    if fill:
        b.fill = fill
    if height:
        ws.row_dimensions[r].height = height
    return b


def est_height(text, per_line=85):
    lines = sum(max(1, len(p) // per_line + 1) for p in str(text or '').split('\n'))
    return max(20, 16 * lines)


def build(rater_file, rater_label):
    items = disagreements(rater_file)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Diskusi'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 95

    ws.cell(row=1, column=1, value='Diskusi: hasil pemeriksaan yang berbeda pendapat').font = FT
    intro = (f'Dari 19 pemeriksaan yang gagal, pendapat Anda ({rater_label}) dan penulis sama pada '
             f'{19 - len(items)} pemeriksaan, dan berbeda pada {len(items)} pemeriksaan di bawah ini. '
             'Tidak ada yang benar atau salah di sini. Tujuannya mencari kesimpulan bersama.\n\n'
             'Untuk setiap nomor: baca penjelasannya, jawab pertanyaan diskusi (boleh sambil mencoba langsung di '
             'browser), lalu isi 2 kotak kuning: KESIMPULAN BERSAMA (pilih dari daftar) dan ALASAN.')
    c = ws.cell(row=2, column=1, value=intro)
    c.font, c.alignment = F, WRAP
    ws.merge_cells('A2:B2')
    ws.row_dimensions[2].height = 95

    dv = DataValidation(type='list', formula1='"' + ','.join(CHOICES) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle='Pilihan tidak ada', error='Pilih dari daftar',
                        promptTitle='Kesimpulan bersama', prompt='Pilih salah satu', showInputMessage=True)
    ws.add_data_validation(dv)

    r = 4
    for n, (item_id, sh, author, rater, comment) in enumerate(items, start=1):
        t = TEXT.get((sh['scenario_id'], sh['project']))
        if t is None:
            sys.exit(f'Belum ada penjelasan bahasa sederhana untuk {sh["scenario_id"]} ({sh["project"]}); tambahkan di TEXT.')
        h = ws.cell(row=r, column=1, value=f'Pemeriksaan nomor {n}')
        h.font, h.fill = FI, HEAD
        ws.cell(row=r, column=2).fill = HEAD
        ws.row_dimensions[r].height = 22
        r += 1
        row(ws, r, 'Kode', item_id)
        r += 1
        row(ws, r, 'Yang dicek', t['dicek'], est_height(t['dicek']))
        r += 1
        row(ws, r, 'Yang terjadi saat diperiksa otomatis', t['terjadi'], est_height(t['terjadi']))
        r += 1
        row(ws, r, 'Pendapat penulis', t['penulis'], est_height(t['penulis']))
        r += 1
        se = f'{PLAIN[rater]}. Catatan Anda: "{comment}"' if comment else PLAIN[rater]
        row(ws, r, f'Pendapat Anda ({rater_label})', se, est_height(se))
        r += 1
        row(ws, r, 'Pertanyaan untuk didiskusikan', t['tanya'], est_height(t['tanya']))
        r += 1
        k = row(ws, r, 'KESIMPULAN BERSAMA', None, 24, INPUT)
        dv.add(k.coordinate)
        r += 1
        row(ws, r, 'ALASAN', None, 60, INPUT)
        r += 2

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f'{OUT.relative_to(RAQA)} ({len(items)} pemeriksaan untuk didiskusikan)')


def do_import(path, rater_file):
    ws = load_workbook(path, data_only=True)['Diskusi']
    items, cur = [], None
    for label, value in ws.iter_rows(min_row=1, max_col=2, values_only=True):
        if label == 'Kode':
            cur = {'item_id': value}
            items.append(cur)
        elif cur is not None and label == 'KESIMPULAN BERSAMA':
            cur['pilihan'] = str(value or '').strip()
        elif cur is not None and label == 'ALASAN':
            cur['alasan'] = ' '.join(str(value or '').split())
    bad = [i['item_id'] for i in items if i.get('pilihan') not in CHOICES]
    if bad:
        sys.exit(f'Kesimpulan bersama kosong/tidak dari daftar untuk: {", ".join(bad)}')
    key = {k['itemId']: k['author'] for k in json.load(open(REVIEW / 'h0-KEY-jangan-dibagikan.json'))}
    print('Kesimpulan bersama:')
    for i in items:
        cat = CHOICES[i['pilihan']]
        print(f"  {i['item_id']}: penulis={key[i['item_id']]} -> FINAL={cat}  | {i.get('alasan') or '(tanpa alasan)'}")
    ov = REVIEW / 'h0-final-overrides.csv'
    with open(ov, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['item_id', 'kategori_final', 'alasan'])
        for i in items:
            w.writerow([i['item_id'], CHOICES[i['pilihan']], i.get('alasan', '')])
    print(f'\nDitulis: {ov.relative_to(RAQA)}')
    print(f'Jalankan: node scripts/h0-failure-classification.js final {Path(rater_file).relative_to(RAQA)} {ov.relative_to(RAQA)}')


if __name__ == '__main__':
    default_rater = str(REVIEW / 'h0-sheet-R2.csv')
    if len(sys.argv) >= 2 and sys.argv[1] == 'build':
        rf = sys.argv[2] if len(sys.argv) > 2 else default_rater
        build(rf, re.search(r'h0-sheet-(\w+)\.csv', rf).group(1))
    elif len(sys.argv) >= 3 and sys.argv[1] == 'import':
        do_import(sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else default_rater)
    else:
        print(__doc__)
